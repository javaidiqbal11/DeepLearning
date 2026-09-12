"""Generate a 16:9 PowerPoint deck for every lecture from ``course_spec``.

Design notes: one idea per slide, a maximum of five bullets, a coloured rule
under every title, and a footer carrying the lecture number. Decks are built
from the spec so the slides can never disagree with the syllabus PDF.
"""
from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt, Emu

sys.path.insert(0, str(Path(__file__).resolve().parent))
from course_spec import COURSE, LECTURES, PARTS, folder_name, slug  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]

# palette ------------------------------------------------------------------
INK = RGBColor(0x1A, 0x1F, 0x2B)
MUTED = RGBColor(0x5A, 0x63, 0x75)
PAPER = RGBColor(0xFF, 0xFF, 0xFF)
WASH = RGBColor(0xF4, 0xF6, 0xFA)
PART_COLOURS = {
    0: RGBColor(0x4A, 0x52, 0x65),   # slate   - ML primer
    1: RGBColor(0x1F, 0x6F, 0xB2),   # blue    - foundations
    2: RGBColor(0x1B, 0x86, 0x6B),   # green   - convolutional vision
    3: RGBColor(0xC1, 0x6A, 0x1C),   # amber   - systems
    4: RGBColor(0x7A, 0x3E, 0xA8),   # violet  - advanced
}

W, H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.72)
BODY_W = W - 2 * MARGIN

FONT = "Calibri"
MONO = "Consolas"


# --------------------------------------------------------------------------
# low-level helpers
# --------------------------------------------------------------------------
def _textbox(slide, left, top, width, height):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    return tf


def _para(tf, text, size, colour=INK, bold=False, space_after=6, first=False,
          align=PP_ALIGN.LEFT, font=FONT, italic=False, space_before=0):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    p.space_before = Pt(space_before)
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = colour
    r.font.name = font
    return p


def _rect(slide, left, top, width, height, colour):
    from pptx.enum.shapes import MSO_SHAPE
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shp.fill.solid()
    shp.fill.fore_color.rgb = colour
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def _blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _footer(slide, lec, text=None):
    tf = _textbox(slide, MARGIN, H - Inches(0.52), BODY_W, Inches(0.3))
    label = text or f"{COURSE['code']}  |  Lecture {lec['number']:02d}  |  {PARTS[lec['part']]}"
    _para(tf, label, 10, MUTED, first=True)


def _title_slide_header(slide, lec, title, kicker=None):
    """Standard content-slide header: kicker, title, coloured rule."""
    accent = PART_COLOURS[lec["part"]]
    top = Inches(0.55)
    if kicker:
        tf = _textbox(slide, MARGIN, top, BODY_W, Inches(0.3))
        _para(tf, kicker.upper(), 11, accent, bold=True, first=True)
        top = top + Inches(0.34)
    tf = _textbox(slide, MARGIN, top, BODY_W, Inches(0.75))
    _para(tf, title, 30, INK, bold=True, first=True)
    rule_top = top + Inches(0.62)
    _rect(slide, MARGIN, rule_top, Inches(1.5), Pt(3.5), accent)
    return rule_top + Inches(0.3)


# --------------------------------------------------------------------------
# slide types
# --------------------------------------------------------------------------
def slide_title(prs, lec):
    s = _blank(prs)
    accent = PART_COLOURS[lec["part"]]
    _rect(s, 0, 0, W, H, WASH)
    _rect(s, 0, 0, Inches(0.34), H, accent)

    tf = _textbox(s, Inches(1.15), Inches(1.5), W - Inches(2.3), Inches(0.4))
    _para(tf, f"{COURSE['code']}  ·  {PARTS[lec['part']]}", 13, accent, bold=True, first=True)

    tf = _textbox(s, Inches(1.15), Inches(2.0), W - Inches(2.3), Inches(0.6))
    _para(tf, f"Lecture {lec['number']:02d}", 17, MUTED, first=True)

    tf = _textbox(s, Inches(1.15), Inches(2.5), W - Inches(2.6), Inches(1.9))
    _para(tf, lec["title"], 40, INK, bold=True, first=True, space_after=10)

    tf = _textbox(s, Inches(1.15), Inches(4.45), W - Inches(2.6), Inches(0.6))
    _para(tf, lec["tagline"], 19, MUTED, italic=True, first=True)

    _rect(s, Inches(1.15), Inches(5.3), Inches(2.0), Pt(3.5), accent)

    tf = _textbox(s, Inches(1.15), Inches(5.7), W - Inches(2.3), Inches(1.0))
    _para(tf, COURSE["title"] + " — " + COURSE["subtitle"], 13, MUTED, first=True, space_after=3)
    _para(tf, COURSE["instructor"], 12, MUTED)
    return s


def slide_objectives(prs, lec):
    s = _blank(prs)
    accent = PART_COLOURS[lec["part"]]
    top = _title_slide_header(s, lec, "Learning objectives", kicker="By the end of this lecture")
    for i, obj in enumerate(lec["objectives"], start=1):
        row_top = top + Inches(0.1) + Inches(1.02) * (i - 1)
        _rect(s, MARGIN, row_top + Inches(0.04), Inches(0.42), Inches(0.42), accent)
        tf = _textbox(s, MARGIN + Inches(0.1), row_top + Inches(0.09), Inches(0.3), Inches(0.35))
        _para(tf, str(i), 15, PAPER, bold=True, first=True, align=PP_ALIGN.CENTER)
        tf = _textbox(s, MARGIN + Inches(0.72), row_top, BODY_W - Inches(0.72), Inches(0.9))
        _para(tf, obj, 17, INK, first=True)
    _footer(s, lec)
    return s


def slide_agenda(prs, lec):
    s = _blank(prs)
    accent = PART_COLOURS[lec["part"]]
    top = _title_slide_header(s, lec, "Agenda", kicker="Where we are going")

    # Pitch adapts to the section count: Lecture 0 has eight, most have five or six.
    n = len(lec["outline"])
    avail = H - top - Inches(1.05)                 # leave room for the lab line + footer
    pitch = min(Inches(0.66), avail / max(n + 1, 1))
    size = 18 if pitch >= Inches(0.6) else 15

    for i, (section, _) in enumerate(lec["outline"], start=1):
        row = top + Inches(0.08) + pitch * (i - 1)
        tf = _textbox(s, MARGIN, row, Inches(0.55), pitch - Inches(0.08))
        _para(tf, f"{i:02d}", size - 1, accent, bold=True, first=True)
        tf = _textbox(s, MARGIN + Inches(0.75), row, BODY_W - Inches(0.75), pitch - Inches(0.08))
        _para(tf, section, size, INK, first=True)
        _rect(s, MARGIN, row + pitch - Inches(0.06), BODY_W, Emu(9525), WASH)

    lab_row = top + Inches(0.08) + pitch * n + Inches(0.08)
    tf = _textbox(s, MARGIN, lab_row, BODY_W, Inches(0.45))
    _para(tf, "Then: hands-on lab  →  " + lec["dataset"], 13, MUTED, italic=True, first=True)
    _footer(s, lec)
    return s


def slide_section_divider(prs, lec, index, section):
    s = _blank(prs)
    accent = PART_COLOURS[lec["part"]]
    _rect(s, 0, 0, W, H, accent)
    tf = _textbox(s, Inches(1.4), Inches(2.85), W - Inches(2.8), Inches(0.5))
    _para(tf, f"SECTION {index:02d}", 14, PAPER, bold=True, first=True)
    tf = _textbox(s, Inches(1.4), Inches(3.3), W - Inches(2.8), Inches(1.4))
    _para(tf, section, 36, PAPER, bold=True, first=True)
    return s


def slide_bullets(prs, lec, section, bullets, cont=False):
    s = _blank(prs)
    accent = PART_COLOURS[lec["part"]]
    title = section + (" (continued)" if cont else "")
    top = _title_slide_header(s, lec, title, kicker=f"Lecture {lec['number']:02d}")

    n = len(bullets)
    size = 19 if n <= 4 else (17 if n == 5 else 16)
    gap = Inches(0.18)
    avail = H - top - Inches(0.85)
    row_h = (avail - gap * (n - 1)) / max(n, 1)

    for i, b in enumerate(bullets):
        row = top + (row_h + gap) * i
        _rect(s, MARGIN, row, BODY_W, row_h, WASH)
        _rect(s, MARGIN, row, Pt(3.5), row_h, accent)
        tf = _textbox(s, MARGIN + Inches(0.28), row + Inches(0.1),
                      BODY_W - Inches(0.5), row_h - Inches(0.16))
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        _para(tf, b, size, INK, first=True)
    _footer(s, lec)
    return s


def slide_lab(prs, lec):
    s = _blank(prs)
    accent = PART_COLOURS[lec["part"]]
    top = _title_slide_header(s, lec, "Hands-on lab", kicker="Notebook")

    _rect(s, MARGIN, top, BODY_W, Inches(1.9), WASH)
    _rect(s, MARGIN, top, Pt(4), Inches(1.9), accent)
    tf = _textbox(s, MARGIN + Inches(0.32), top + Inches(0.22),
                  BODY_W - Inches(0.64), Inches(1.5))
    _para(tf, lec["lab"], 17, INK, first=True)

    y = top + Inches(2.2)
    rows = [
        ("Notebook", f"notebooks/L{lec['number']:02d}_{slug(lec['title'])[:44]}.ipynb"),
        ("Data", lec["dataset"]),
        ("Helpers", "from dlcourse import load_shapes, train, evaluate, show_grid"),
    ]
    if lec.get("app"):
        rows.append(("Service", lec["app"]))
    for label, value in rows:
        tf = _textbox(s, MARGIN, y, Inches(1.5), Inches(0.4))
        _para(tf, label, 14, accent, bold=True, first=True)
        tf = _textbox(s, MARGIN + Inches(1.6), y, BODY_W - Inches(1.6), Inches(0.45))
        _para(tf, value, 14, INK, first=True, font=MONO)
        y = y + Inches(0.52)
    _footer(s, lec)
    return s


def slide_tasks(prs, lec):
    s = _blank(prs)
    accent = PART_COLOURS[lec["part"]]
    core = [t for t in lec["tasks"] if t[0] == "core"]
    stretch = [t for t in lec["tasks"] if t[0] == "stretch"]
    top = _title_slide_header(s, lec, "Your tasks", kicker="Due before the next lecture")

    tf = _textbox(s, MARGIN, top, Inches(6.0), Inches(0.35))
    _para(tf, f"CORE — all {len(core)} required", 13, accent, bold=True, first=True)
    y = top + Inches(0.42)
    for i, (_, title, _desc) in enumerate(core, start=1):
        tf = _textbox(s, MARGIN, y, Inches(6.0), Inches(0.42))
        _para(tf, f"{i}.  {title}", 15, INK, first=True)
        y = y + Inches(0.45)

    x2 = MARGIN + Inches(6.5)
    tf = _textbox(s, x2, top, Inches(5.4), Inches(0.35))
    _para(tf, f"STRETCH — {len(stretch)} optional", 13, MUTED, bold=True, first=True)
    y2 = top + Inches(0.42)
    for i, (_, title, _desc) in enumerate(stretch, start=1):
        tf = _textbox(s, x2, y2, Inches(5.4), Inches(0.42))
        _para(tf, f"{i}.  {title}", 15, MUTED, first=True)
        y2 = y2 + Inches(0.45)

    tf = _textbox(s, MARGIN, H - Inches(1.05), BODY_W, Inches(0.45))
    _para(tf, f"Full instructions: tasks/Lecture_{lec['number']:02d}_Tasks.md", 13, MUTED,
          italic=True, first=True)
    _footer(s, lec)
    return s


def slide_reading(prs, lec):
    s = _blank(prs)
    accent = PART_COLOURS[lec["part"]]
    top = _title_slide_header(s, lec, "Reading", kicker="Go deeper")
    for i, r in enumerate(lec["reading"], start=1):
        row = top + Inches(0.62) * i - Inches(0.5)
        tf = _textbox(s, MARGIN, row, Inches(0.4), Inches(0.5))
        _para(tf, "▸", 16, accent, bold=True, first=True)
        tf = _textbox(s, MARGIN + Inches(0.45), row, BODY_W - Inches(0.45), Inches(0.6))
        _para(tf, r.replace("*", ""), 16, INK, first=True)
    _footer(s, lec)
    return s


def slide_summary(prs, lec):
    s = _blank(prs)
    accent = PART_COLOURS[lec["part"]]
    _rect(s, 0, 0, W, H, WASH)
    top = _title_slide_header(s, lec, "What to take away", kicker="Summary")
    # The first bullet of each section is its headline claim.
    points = [bullets[0] for _section, bullets in lec["outline"]]
    n = len(points)
    avail = H - top - Inches(1.25)                 # room for the "next lecture" line
    pitch = min(Inches(0.78), avail / max(n, 1))
    size = 17 if pitch >= Inches(0.75) else (16 if pitch >= Inches(0.62) else 14)

    for i, p in enumerate(points):
        row = top + Inches(0.04) + pitch * i
        _rect(s, MARGIN, row + Inches(0.16), Inches(0.14), Inches(0.14), accent)
        tf = _textbox(s, MARGIN + Inches(0.40), row, BODY_W - Inches(0.40), pitch - Inches(0.06))
        _para(tf, p, size, INK, first=True)

    nxt = lec["number"] + 1
    if nxt <= 16:
        nxt_lec = LECTURES[nxt]
        tf = _textbox(s, MARGIN, H - Inches(1.1), BODY_W, Inches(0.5))
        _para(tf, f"Next — Lecture {nxt:02d}: {nxt_lec['title']}", 14, accent, bold=True, first=True)
    _footer(s, lec)
    return s


# --------------------------------------------------------------------------
def build_deck(lec) -> Presentation:
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H

    slide_title(prs, lec)
    slide_objectives(prs, lec)
    slide_agenda(prs, lec)

    for i, (section, bullets) in enumerate(lec["outline"], start=1):
        slide_section_divider(prs, lec, i, section)
        # Never more than 5 bullets on one slide.
        chunks = [bullets[j:j + 5] for j in range(0, len(bullets), 5)] or [[]]
        for c, chunk in enumerate(chunks):
            slide_bullets(prs, lec, section, chunk, cont=(c > 0))

    slide_lab(prs, lec)
    slide_tasks(prs, lec)
    slide_reading(prs, lec)
    slide_summary(prs, lec)
    return prs


def main():
    total = 0
    for lec in LECTURES:
        out_dir = ROOT / "lectures" / folder_name(lec["number"]) / "slides"
        out_dir.mkdir(parents=True, exist_ok=True)
        path = out_dir / f"Lecture_{lec['number']:02d}_{slug(lec['title'])}.pptx"
        prs = build_deck(lec)
        prs.save(path)
        n = len(prs.slides.__iter__.__self__._sldIdLst)
        total += n
        print(f"  L{lec['number']:02d}  {n:2d} slides  {path.name}")
    print(f"Built {len(LECTURES)} decks, {total} slides total.")


if __name__ == "__main__":
    main()
